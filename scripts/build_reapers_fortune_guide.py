from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "Reapers_Fortune_Game_Guide.docx"
ASSETS = ROOT / "public" / "assets" / "gems"

# compact_reference_guide token map, with named brand-color overrides.
PAGE_W = 12240
PAGE_H = 15840
MARGIN = 1440
CONTENT_W = 9360
TABLE_INDENT = 120
CELL_TOP_BOTTOM = 80
CELL_START_END = 120
FONT = "Calibri"
INK = "24152D"
MUTED = "6F6374"
BLUE = "2E74B5"
BLUE_DARK = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
MAGENTA = "B82672"  # named override: Reaper brand accent
PALE_MAGENTA = "F9E8F1"  # named override: Reaper callout fill
GOLD = "8A6510"  # named override: jackpot/caution accent
RED = "9B1C1C"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=CELL_TOP_BOTTOM, start=CELL_START_END,
                     bottom=CELL_TOP_BOTTOM, end=CELL_START_END):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=TABLE_INDENT):
    assert sum(widths) == CONTENT_W, (widths, sum(widths))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_W))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(width))
        grid.append(gc)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths[idx] / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_repeat_keep(paragraph, keep_next=False, keep_lines=True):
    p_pr = paragraph._p.get_or_add_pPr()
    if keep_next:
        el = OxmlElement("w:keepNext")
        p_pr.append(el)
    if keep_lines:
        el = OxmlElement("w:keepLines")
        p_pr.append(el)


def set_run_font(run, size=None, color=INK, bold=None, italic=None, name=FONT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def style_paragraph(p, before=0, after=6, line=1.25, align=None):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if align is not None:
        p.alignment = align


def add_text(doc, text, bold_lead=None, after=6, color=INK):
    p = doc.add_paragraph()
    style_paragraph(p, after=after)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size=11, color=color, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, size=11, color=color)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11, color=color)
    return p


def add_callout(doc, label, text, fill=PALE_MAGENTA, accent=MAGENTA, trailing_spacer=True):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_W])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    style_paragraph(p, after=0)
    r = p.add_run(label + "  ")
    set_run_font(r, size=10.5, color=accent, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=INK)
    if trailing_spacer:
        spacer = doc.add_paragraph()
        style_paragraph(spacer, after=3)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    set_repeat_keep(p, keep_next=True)
    return p


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    values = {
        "Heading 1": (16, MAGENTA, 18, 10),
        "Heading 2": (13, MAGENTA, 14, 7),
        "Heading 3": (12, BLUE_DARK, 10, 5),
    }
    for name, (size, color, before, after) in values.items():
        st = styles[name]
        st.font.name = FONT
        st._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        st._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.0
        st.paragraph_format.keep_with_next = True


def create_numbering(doc):
    numbering = doc.part.numbering_part.element

    def make_num(abstract_id, num_id, fmt, text, left=540, hanging=270):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        lvl.append(lvl_text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        lvl.append(suff)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), str(left))
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str(left))
        ind.set(qn("w:hanging"), str(hanging))
        p_pr.append(ind)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.append(spacing)
        lvl.append(p_pr)
        r_pr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), FONT)
        fonts.set(qn("w:hAnsi"), FONT)
        r_pr.append(fonts)
        lvl.append(r_pr)
        abstract.append(lvl)
        numbering.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        aid = OxmlElement("w:abstractNumId")
        aid.set(qn("w:val"), str(abstract_id))
        num.append(aid)
        numbering.append(num)

    make_num(100, 100, "decimal", "%1.")
    make_num(101, 101, "bullet", "•")
    return 100, 101


def add_list_item(doc, text, num_id, bold_lead=None):
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    p_pr.append(num_pr)
    style_paragraph(p, after=4)
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run_font(r, size=11, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_run_font(r, size=11)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11)
    return p


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color=MUTED)


def setup_page(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)
    sec.different_first_page_header_footer = True
    doc.settings.odd_and_even_pages_header_footer = False

    for header in (sec.header, sec.even_page_header):
        h = header.paragraphs[0]
        h.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style_paragraph(h, after=0)
        r = h.add_run("REAPER'S FORTUNE  |  GAME GUIDE")
        set_run_font(r, size=8.5, color=MUTED, bold=True)
    for footer in (sec.footer, sec.even_page_footer):
        f = footer.paragraphs[0]
        f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        style_paragraph(f, after=0)
        r = f.add_run("Page ")
        set_run_font(r, size=9, color=MUTED)
        add_page_number(f)


def set_cell_text(cell, text, bold=False, color=INK, size=10, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell.paragraphs[0]
    p.alignment = align
    style_paragraph(p, after=0, line=1.15)
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=bold)


def add_cover(doc):
    p = doc.add_paragraph()
    style_paragraph(p, after=18, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("PLAYER & IMPLEMENTATION GUIDE")
    set_run_font(r, size=10, color=MAGENTA, bold=True)

    portrait = ROOT / "assets" / "character" / "reaper-queen-portrait.png"
    if portrait.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, after=12)
        p.add_run().add_picture(str(portrait), height=Inches(3.6))

    p = doc.add_paragraph()
    style_paragraph(p, after=7, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("REAPER'S FORTUNE")
    set_run_font(r, size=29, color=INK, bold=True)

    p = doc.add_paragraph()
    style_paragraph(p, after=18, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("How the demo works, what counts as a match, and how prizes are selected")
    set_run_font(r, size=13, color=MUTED)

    add_callout(
        doc,
        "THE SIMPLE RULE",
        "A visible win is 3 or more matching symbols on consecutive reels, starting at Reel 1 and following one payline. A Wild may replace the matching symbol.",
    )

    p = doc.add_paragraph()
    style_paragraph(p, before=20, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("Based on index.html  |  3 September 2026")
    set_run_font(r, size=9.5, color=MUTED, italic=True)
    doc.add_page_break()


def add_spin_flow(doc, decimal_id, bullet_id):
    add_heading(doc, "1. What the player does", 1)
    add_text(doc, "Reaper's Fortune is a 5-reel, 3-row slot-style demo. The player does not aim or time the reels; each spin is random.")
    steps = [
        "Choose a bet: 10, 20, 30, or 50 credits.",
        "Press SPIN. The selected bet is deducted immediately.",
        "The game randomly chooses a payout multiplier.",
        "The reels animate and stop one at a time.",
        "If the chosen multiplier is above zero, the prize is added to the balance.",
    ]
    for step in steps:
        add_list_item(doc, step, decimal_id)

    add_callout(doc, "PRIZE FORMULA", "Prize credited = selected bet x payout multiplier.", fill=CALLOUT, accent=BLUE_DARK)

    add_heading(doc, "2. What counts as a visible win", 1)
    rules = [
        ("Start on Reel 1. ", "A matching group beginning on Reel 2, 3, 4, or 5 does not qualify."),
        ("Use consecutive reels. ", "The match cannot skip a reel."),
        ("Follow one payline. ", "Symbols do not need to be in a straight row if the selected payline bends."),
        ("Match at least 3 symbols. ", "Three, four, or five matching positions form the visible winning sequence."),
        ("Wild substitutes. ", "A Wild can stand in for the regular matching symbol in the generated win."),
    ]
    for lead, body in rules:
        add_list_item(doc, lead + body, bullet_id, bold_lead=lead)

    add_text(doc, "A Scatter symbol currently has no bonus or payout behavior. Despite the game's subtitle, there is no implemented free-spin or scatter bonus.", bold_lead="A Scatter symbol currently has no bonus or payout behavior. ")

    doc.add_page_break()
    add_heading(doc, "3. The game symbols", 1)
    add_text(doc, "Symbol weights control only the random filler grid. They do not set the prize; the selected payout multiplier does that.")
    symbols = [
        ("blood-heart.png", "Blood Heart", "20%", "Regular matching symbol"),
        ("chained-soulstone.png", "Chained Soulstone", "18%", "Regular matching symbol"),
        ("amethyst-skull.png", "Amethyst Skull", "18%", "Regular matching symbol"),
        ("soul-orb.png", "Soul Orb", "17%", "Regular matching symbol"),
        ("seven.png", "Seven", "12%", "High-tier matching symbol"),
        ("red.png", "Wild", "8%", "May substitute for a matching symbol"),
        ("green.png", "Scatter", "7%", "Visual filler only; no current feature"),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    set_table_geometry(table, [950, 2500, 1000, 4910])
    for i, text in enumerate(("Symbol", "Name", "Weight", "Current role")):
        set_cell_shading(table.rows[0].cells[i], LIGHT_BLUE)
        set_cell_text(table.rows[0].cells[i], text, bold=True, color=BLUE_DARK, align=WD_ALIGN_PARAGRAPH.CENTER if i in (0, 2) else WD_ALIGN_PARAGRAPH.LEFT)
    set_repeat_table_header(table.rows[0])
    for filename, name, weight, role in symbols:
        cells = table.add_row().cells
        set_table_geometry(table, [950, 2500, 1000, 4910])
        pic_p = cells[0].paragraphs[0]
        pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(pic_p, after=0)
        image = ASSETS / filename
        if image.exists():
            pic_p.add_run().add_picture(str(image), height=Inches(0.40))
        set_cell_text(cells[1], name, bold=True)
        set_cell_text(cells[2], weight, bold=True, color=BLUE_DARK, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[3], role)
    add_callout(doc, "REMEMBER", "The game may replace selected positions after this filler grid is generated in order to display the predetermined payout.", fill=CALLOUT, accent=BLUE_DARK)
    doc.add_page_break()


def add_paylines_and_examples(doc):
    add_heading(doc, "4. The 10 paylines", 1)
    add_text(doc, "Read every pattern from Reel 1 to Reel 5. T = top row, M = middle row, and B = bottom row.")
    lines = [
        ("1", "T - T - T - T - T", "Top horizontal"),
        ("2", "M - M - M - M - M", "Middle horizontal"),
        ("3", "B - B - B - B - B", "Bottom horizontal"),
        ("4", "T - M - B - M - T", "Down and up V"),
        ("5", "B - M - T - M - B", "Up and down inverted V"),
        ("6", "T - T - M - B - B", "Stepped downward"),
        ("7", "B - B - M - T - T", "Stepped upward"),
        ("8", "M - T - T - T - M", "Upper arch"),
        ("9", "M - B - B - B - M", "Lower arch"),
        ("10", "T - M - M - M - T", "Shallow V"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_table_geometry(table, [900, 3900, 4560])
    for i, text in enumerate(("Line", "Row pattern", "Shape")):
        set_cell_shading(table.rows[0].cells[i], LIGHT_BLUE)
        set_cell_text(table.rows[0].cells[i], text, bold=True, color=BLUE_DARK, align=WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT)
    set_repeat_table_header(table.rows[0])
    for n, pattern, shape in lines:
        cells = table.add_row().cells
        set_table_geometry(table, [900, 3900, 4560])
        set_cell_text(cells[0], n, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[1], pattern, bold=True, color=MAGENTA)
        set_cell_text(cells[2], shape)

    add_heading(doc, "5. Worked match examples", 1)
    add_heading(doc, "Example A: a valid 3-match using a Wild", 2)
    add_text(doc, "On the middle horizontal payline, Amethyst + Amethyst + Wild is treated visually as three Amethysts. The Blood Heart on Reel 4 stops the sequence, so this is a 3-match.")
    add_grid_example(doc, [
        ["Soul Orb", "Chain", "Chain", "Blood", "Blood"],
        ["Amethyst", "Amethyst", "WILD", "Blood", "Scatter"],
        ["Chain", "Amethyst", "Chain", "Blood", "Soul Orb"],
    ], highlight={(1, 0), (1, 1), (1, 2)})

    doc.add_page_break()
    add_heading(doc, "Example B: not a win", 2)
    add_text(doc, "Three identical symbols that begin on Reel 2 do not qualify because every match must start on Reel 1.")
    add_grid_example(doc, [
        ["Soul Orb", "Seven", "Seven", "Seven", "Blood"],
        ["Amethyst", "Chain", "Scatter", "Blood", "Soul Orb"],
        ["Chain", "Amethyst", "Soul Orb", "Blood", "Scatter"],
    ], highlight=set())


def add_grid_example(doc, rows, highlight):
    table = doc.add_table(rows=4, cols=5)
    table.style = "Table Grid"
    set_table_geometry(table, [1872] * 5)
    for c in range(5):
        set_cell_shading(table.rows[0].cells[c], INK)
        set_cell_text(table.rows[0].cells[c], f"REEL {c + 1}", bold=True, color=WHITE, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    for r in range(3):
        for c in range(5):
            cell = table.rows[r + 1].cells[c]
            if (r, c) in highlight:
                set_cell_shading(cell, PALE_MAGENTA)
            set_cell_text(cell, rows[r][c], bold=(r, c) in highlight, color=MAGENTA if (r, c) in highlight else INK, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                set_repeat_keep(p, keep_next=row_idx < len(table.rows) - 1)
    spacer = doc.add_paragraph()
    style_paragraph(spacer, after=2)


def add_payout_model(doc):
    add_heading(doc, "6. Current payout model", 1)
    add_callout(doc, "IMPORTANT", "The demo chooses the payout first and then builds a reel picture that illustrates it. It does not scan the completed grid to calculate the prize.", fill=PALE_MAGENTA, accent=RED)
    add_text(doc, "The configured outcome table is:")
    rows = [
        ("64%", "x0", "Lose the full bet", "Lose"),
        ("16%", "x0.5", "Receive half the bet back", "Small"),
        ("10%", "x1", "Receive the full bet back", "Small"),
        ("5%", "x2", "Receive twice the bet", "Win"),
        ("3%", "x4", "Receive four times the bet", "Win"),
        ("1.5%", "x10", "Receive ten times the bet", "Big"),
        ("0.5%", "x50", "Receive fifty times the bet", "Jackpot"),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    set_table_geometry(table, [1500, 1400, 4360, 2100])
    for i, text in enumerate(("Chance", "Multiplier", "Amount credited", "Tier")):
        set_cell_shading(table.rows[0].cells[i], LIGHT_BLUE)
        set_cell_text(table.rows[0].cells[i], text, bold=True, color=BLUE_DARK, align=WD_ALIGN_PARAGRAPH.CENTER if i != 2 else WD_ALIGN_PARAGRAPH.LEFT)
    set_repeat_table_header(table.rows[0])
    for chance, mult, result, tier in rows:
        cells = table.add_row().cells
        set_table_geometry(table, [1500, 1400, 4360, 2100])
        set_cell_text(cells[0], chance, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[1], mult, bold=True, color=MAGENTA, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[2], result)
        set_cell_text(cells[3], tier, align=WD_ALIGN_PARAGRAPH.CENTER)

    add_heading(doc, "How many symbols the demo inserts", 2)
    insertion = [
        ("x0", "No forced match; random filler grid"),
        ("x0.5, x1, x2", "3 matching positions on one payline"),
        ("x4", "4 matching positions on one payline"),
        ("x10", "5 matching positions on one payline"),
        ("x50", "5 Wild symbols on one payline"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [2400, 6960])
    for i, text in enumerate(("Selected payout", "Reel presentation")):
        set_cell_shading(table.rows[0].cells[i], LIGHT_GRAY)
        set_cell_text(table.rows[0].cells[i], text, bold=True, color=BLUE_DARK)
    set_repeat_table_header(table.rows[0])
    for mult, presentation in insertion:
        cells = table.add_row().cells
        set_table_geometry(table, [2400, 6960])
        set_cell_text(cells[0], mult, bold=True, color=MAGENTA, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[1], presentation)

    add_heading(doc, "RTP calculation", 2)
    add_text(doc, "Although the interface says 90.00% RTP, the configured probabilities return 80.00% on average:")
    p = doc.add_paragraph()
    style_paragraph(p, after=7, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("(16% x 0.5) + (10% x 1) + (5% x 2) + (3% x 4) + (1.5% x 10) + (0.5% x 50) = 80%")
    set_run_font(r, size=10.5, color=BLUE_DARK, bold=True)
    add_callout(doc, "EXPECTED RESULT", "The long-run expected loss is 20% of each wager. This conflicts with the 90% label shown in the interface.", fill=CALLOUT, accent=RED)
def add_known_issues(doc, bullet_id):
    add_heading(doc, "7. Important implementation notes", 1)
    add_heading(doc, "Five Wilds do not uniquely identify the jackpot", 2)
    add_text(doc, "The x10 tier randomly chooses Seven, Wild, or Soul Orb as its five-symbol match. If it chooses Wild, the screen shows five Wilds but the award remains x10. The x50 tier also shows five Wilds.")
    add_callout(doc, "CURRENT BEHAVIOR", "x50 always displays five Wilds, but a display of five Wilds does not always mean x50.", fill=PALE_MAGENTA, accent=RED)

    add_heading(doc, "Accidental visible matches are not evaluated", 2)
    add_text(doc, "The random filler grid can occasionally form an additional match. Because the game never scans the finished grid, that accidental match is not paid or highlighted. Only the positions inserted for the selected outcome are treated as winners.")

    add_heading(doc, "Wild and Scatter are presentation features", 2)
    notes = [
        "Wild substitution is not implemented through a general line evaluator. Wilds are inserted while constructing a predetermined winning line.",
        "Scatter has a 7% filler weight but no scatter count, bonus, or payout rule.",
        "There are no symbol-specific pay values. The selected multiplier alone determines the prize.",
    ]
    for note in notes:
        add_list_item(doc, note, bullet_id)

    doc.add_page_break()
    add_heading(doc, "8. Controls and session behavior", 1)
    controls = [
        ("Balance", "Starts at 10,000.00 and resets when the page reloads."),
        ("Last win", "Resets to 0 at the start of every spin and then shows the most recent credited prize."),
        ("Auto x10", "Starts a ten-spin sequence. After a completed spin it stops if the next bet is unaffordable; if started with insufficient balance, the button can remain active due to a demo edge case."),
        ("Sound", "The button changes between on and off, but the demo contains no audio playback."),
        ("Storage", "No balance or history is saved locally or sent to a server."),
        ("Randomness", "Uses the browser's Math.random(); it is not a certified gambling random-number generator."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [2700, 6660])
    for i, text in enumerate(("Feature", "Current behavior")):
        set_cell_shading(table.rows[0].cells[i], LIGHT_BLUE)
        set_cell_text(table.rows[0].cells[i], text, bold=True, color=BLUE_DARK)
    set_repeat_table_header(table.rows[0])
    for label, behavior in controls:
        cells = table.add_row().cells
        set_table_geometry(table, [2700, 6660])
        set_cell_text(cells[0], label, bold=True, color=MAGENTA)
        set_cell_text(cells[1], behavior)
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                set_repeat_keep(p, keep_next=row_idx < len(table.rows) - 1)

    add_callout(doc, "SUMMARY", "The player chooses a bet and spins; the demo randomly chooses a multiplier, builds a matching reel pattern to represent that multiplier, and credits bet x multiplier to the balance.", fill=CALLOUT, accent=BLUE_DARK, trailing_spacer=False)


def build():
    doc = Document()
    setup_page(doc)
    setup_styles(doc)
    decimal_id, bullet_id = create_numbering(doc)
    add_cover(doc)
    add_spin_flow(doc, decimal_id, bullet_id)
    add_paylines_and_examples(doc)
    add_payout_model(doc)
    add_known_issues(doc, bullet_id)

    props = doc.core_properties
    props.title = "Reaper's Fortune - Game Guide"
    props.subject = "Player rules and implementation behavior"
    props.author = "Soul Hunter Project"
    props.keywords = "Reaper's Fortune, slot demo, paylines, Wild, RTP"
    props.comments = "Generated from the current HTML demo behavior."
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
